#!/usr/bin/env python3
"""개발완료보고서 — Word(.docx) 판. PPT 판(build_pptx.py)과 같은 내용·수치를 글과 표로 풀어 쓴다.

영상은 시간 순서의 장면 캡처 4컷으로 넣는다. 수치는 PPT 판과 근거 문서(sw/docs/evidence)와 같다.
사용: python3 sw/report/build_report_docx.py   (python-docx, Pillow, ffmpeg 필요)
출력: sw/report/2026ESWContest_모빌리티_ARTS_개발완료보고서.docx
"""
from __future__ import annotations

import pathlib
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parents[2]
HERE = pathlib.Path(__file__).resolve().parent
A = HERE / "assets"; V = A / "video"; FIG = ROOT / "sw/docs/figures"; HW = ROOT / "hw/images"
CAPS = A / "word"; CAPS.mkdir(exist_ok=True)
OUT = HERE / "2026ESWContest_모빌리티_ARTS_개발완료보고서.docx"   # 제출: 이 문서를 PDF 로 변환 (to_pdf_word.sh)
REPO = "https://github.com/gkgk0119gmail-arch/2026ESWContest_mobility_ARTS"
FONT = "나눔고딕"; BLUE = RGBColor(0x00, 0x2C, 0x5F); RED = RGBColor(0xC0, 0x2B, 0x2B); GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)
TEXT_W = 17.0


def _font(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(k), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color


st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.35; st.paragraph_format.space_after = Pt(4)
_fig_n = [0]; _tab_n = [0]


def P(text="", size=10, bold=False, color=None, align=None, after=4, before=0, italic=False):
    """**굵게** 표시를 지원하는 문단."""
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part: continue
        r = p.add_run(part); _font(r, size, bold or (i % 2 == 1), color); r.italic = italic
    return p


def B(text, level=0):
    p = P("· " + text, after=2); p.paragraph_format.left_indent = Cm(0.4 + 0.5 * level); p.paragraph_format.first_line_indent = Cm(-0.3)
    return p


def H1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); _font(r, 16, True, BLUE)
    b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "2"), ("w:color", "002C5F")): bt.set(qn(k), v)
    b.append(bt); p._p.get_or_add_pPr().append(b)
    return p


def H2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); _font(r, 12, True, BLUE); return p


def FIGURE(path, width_cm, caption):
    _fig_n[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    P(f"그림 {_fig_n[0]}. {caption}", size=8.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor); tcPr.append(s)


def TABLE(caption, head, rows, widths_cm, size=9):
    _tab_n[0] += 1
    P(f"표 {_tab_n[0]}. {caption}", size=9, bold=True, color=BLUE, after=2, before=4).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(head)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False                                               # 지정한 칸 너비를 지키게 (Word·LibreOffice 둘 다)
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, wcm in zip(grid.findall(qn("w:gridCol")), widths_cm): gc.set(qn("w:w"), str(int(wcm * 567)))
    for ri, row in enumerate([head] + rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.width = Cm(widths_cm[ci]); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
            parts = str(val).split("**")
            for i, part in enumerate(parts):
                if part: _font(p.add_run(part), size, ri == 0 or i % 2 == 1, RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
            if ri == 0: _shade(c, "002C5F")
            elif ci == 0: _shade(c, "EEF2F7")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def filmstrip(mp4, name, k=4):
    """영상에서 시간 순서로 4 장면을 뽑아 2×2 로 붙이고 번호·시각을 단다 (PDF·Word 에서는 영상이 재생되지 않는다)."""
    from PIL import Image, ImageDraw, ImageFont
    out = CAPS / f"{name}.jpg"
    if out.exists() and out.stat().st_mtime > pathlib.Path(mp4).stat().st_mtime: return out
    dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(mp4)],
                               capture_output=True, text=True).stdout.strip())
    ts = [dur * f for f in (0.12, 0.40, 0.68, 0.94)][:k]; frames = []
    for i, t_ in enumerate(ts):
        f = CAPS / f"_{name}_{i}.jpg"
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-y", "-ss", f"{t_:.2f}", "-i", str(mp4), "-frames:v", "1", str(f)], check=True)
        frames.append(Image.open(f).convert("RGB")); f.unlink()
    fw, fh = frames[0].size; g = 8
    canvas = Image.new("RGB", (fw * 2 + g, fh * 2 + g), (255, 255, 255)); d = ImageDraw.Draw(canvas)
    fnt = ImageFont.truetype(str(HERE / "fonts/Pretendard-Bold.otf"), max(18, fh // 18))
    for i, im in enumerate(frames):
        x0, y0 = (i % 2) * (fw + g), (i // 2) * (fh + g); canvas.paste(im, (x0, y0))
        lab = f" {i + 1}   {ts[i]:.1f} s "; tw = d.textlength(lab, font=fnt)
        d.rectangle((x0, y0, x0 + tw + 8, y0 + fh // 14 + 10), fill=(0, 44, 95)); d.text((x0 + 4, y0 + 3), lab, font=fnt, fill=(255, 255, 255))
    canvas.save(out, quality=90); return out


def grid(paths, name, cols=None, cell_h=360, labels=None):
    """그림 여러 장을 같은 높이로 맞춰 한 줄(또는 여러 줄)로 붙인다. labels 가 있으면 왼쪽 아래에 이름표."""
    from PIL import Image, ImageDraw, ImageFont
    out = CAPS / f"{name}.jpg"
    ims = [Image.open(p).convert("RGB") for p in paths]
    ims = [im.resize((int(im.width * cell_h / im.height), cell_h), Image.LANCZOS) for im in ims]
    cols = cols or len(ims); g = 10
    rows = [ims[i:i + cols] for i in range(0, len(ims), cols)]
    W = max(sum(im.width for im in r) + g * (len(r) - 1) for r in rows); Hh = len(rows) * cell_h + g * (len(rows) - 1)
    canvas = Image.new("RGB", (W, Hh), (255, 255, 255)); d = ImageDraw.Draw(canvas)
    fnt = ImageFont.truetype(str(HERE / "fonts/Pretendard-Bold.otf"), max(18, cell_h // 16))
    k = 0
    for ri, r in enumerate(rows):
        x = 0
        for im in r:
            canvas.paste(im, (x, ri * (cell_h + g)))
            if labels:
                lab = f" {labels[k]} "; tw = d.textlength(lab, font=fnt); y0 = ri * (cell_h + g) + cell_h - cell_h // 11 - 8
                d.rectangle((x, y0, x + tw + 6, y0 + cell_h // 11 + 8), fill=(0, 44, 95)); d.text((x + 3, y0 + 2), lab, font=fnt, fill=(255, 255, 255))
            x += im.width + g; k += 1
    canvas.save(out, quality=90); return out


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def page_number_footer(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IcePredict · 팀 ARTS · 제24회 임베디드SW경진대회   —   "); _font(r, 8, color=GREY)
    r = p.add_run(); _font(r, 8, color=GREY)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)


# ═════════════════════════════ 표지 ═════════════════════════════
P("", after=60)
P("제24회 임베디드SW경진대회 · 모빌리티 부문", size=12, color=BLUE, bold=True)
P("개발완료보고서", size=14, color=GREY, after=18)
P("IcePredict", size=40, bold=True, color=BLUE, after=4)
P("AI 예측과 RTOS 반응의 이중 안전망 블랙아이스 대응 시스템", size=17, bold=True, after=8)
P("카메라가 블랙아이스를 미리 보고, 놓치면 노면 마찰을 직접 재서 제때 세운다.", size=12, color=GREY, after=30)
pc = doc.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
pc.add_run().add_picture(str(A / "intro_work_photos.png"), width=Cm(TEXT_W))
P("", after=30)
TABLE("팀 정보", ["항목", "내용"],
      [["팀명", "ARTS"], ["팀원", "남윤상(팀장) · 김진찬 · 이지성"], ["작품명", "IcePredict"],
       ["주요 구성", "STM32N6570-DK(Neural-ART NPU · ThreadX) · Raspberry Pi 5 + AI HAT+ 2 · CARLA HIL · Moza 휠·페달"],
       ["소스 코드·근거", REPO]], [3.2, 13.8], size=9.5)
page_number_footer(sec)
page_break()

# ═════════════════════════════ 목차 ═════════════════════════════
H1("목차")
for line in ["1. 개요", "2. 문제 정의 — 왜 블랙아이스는 카메라만으로 풀리지 않는가", "3. 시스템 구성", "4. 개발계획서 대비 변경점",
             "5. 1차 방어 — 카메라가 미리 본다", "6. 2차 방어 — 카메라가 놓치면 차량 거동으로 받는다",
             "7. 2차 방어 확장 — 노면 마찰을 재서 제동 시점을 고친다", "8. 안전 전용 RTOS — 메인 컴퓨터가 멈춰도 2차 방어는 산다",
             "9. 시연 구성", "10. 검증 방법과 재현성", "11. 한계와 향후 계획", "부록. 근거 문서 목록"]:
    P(line, size=11, after=3)
page_break()

# ═════════════════════════════ 1. 개요 ═════════════════════════════
H1("1. 개요")
P("IcePredict 는 블랙아이스 사고를 줄이는 **운전자 보조 장치**다. 조향은 운전자에게 두고 경고·가속 제한·제동만 돕는다. "
  "두 겹의 방어로 구성된다. **1차 방어**는 카메라 AI 가 노면을 보고 빙판에 닿기 전에 경고하고 감속한다. "
  "**2차 방어**는 카메라가 놓쳤을 때 IMU 로 차량 거동을 보고 미끄러짐을 확정하며, 확장된 2차 방어는 노면 마찰을 직접 재서 "
  "'언제부터 서야 하는지'를 노면에 맞게 다시 계산한다. 2차 방어 판단은 메인 컴퓨터(리눅스)와 분리된 **안전 전용 RTOS** 에서 돌아 "
  "메인 컴퓨터가 멈춰도 계속된다.")
TABLE("핵심 결과", ["구분", "결과", "조건"],
      [["1차 방어", "블랙아이스 경보율 **96.6 %**, 마른 노면 오경보율 **0.3 %**", "실제 도로 사진 69,358장을 STM32N6 보드 NPU 에 직접 입력"],
       ["2차 방어 응답", "최악 **22.6 µs** (리눅스 최악 5,790 µs)", "STM32N6 + ThreadX, 49,405 표본"],
       ["2차 방어 효과", "방어 없음 9/9 제어 상실 → 2차 방어 차선 안 정지", "CARLA HIL, 같은 빙판·날씨·주행선"],
       ["마찰 인식 제동", "정지 **2/9 → 8/9** (7회는 목표대로 앞차 4 m 앞)", "CARLA, 빙판 위 정지 차량, 마찰 3종 × 속도 3종"],
       ["고장 격리", "메인 컴퓨터 정지 시 충돌 속도 **53 → 39 km/h**", "곡선 빙판 μ 0.12 · 60 km/h"],
       ["실시간성", "전 코어 부하에서 20 ms 주기 6,000회 마감 놓침 **0** (최악 0.06 ms)", "파이 격리 코어(안전 전용 RTOS 대체)"]],
      [2.8, 7.6, 6.6])
FIGURE(filmstrip(V / "v_compare_bev.mp4", "compare_bev"), TEXT_W,
       "같은 조건 비교(조감) — 각 칸의 왼쪽은 방어 없음(제어를 잃고 미끄러짐), 오른쪽은 STM32N6 RTOS 2차 방어(차선 안 정지). 시간 순서 4장면")

# ═════════════════════════════ 2. 문제 ═════════════════════════════
H1("2. 문제 정의 — 왜 블랙아이스는 카메라만으로 풀리지 않는가")
B("**라벨이 없다.** 공개 데이터셋에 '블랙아이스' 라벨이 없다. RSCD 의 ice 57,262장도 다져진 눈·서리에 가까운 '얼음 노면'일 뿐, 투명한 블랙아이스를 따로 구분하지 않는다.")
B("**원래 어려운 문제다.** 투명·검은 얼음은 아스팔트와 시각적으로 구분되지 않는다. 우리 모델도 실사진 얼음 6,340장 중 215장(3.4 %)을 어떤 문턱으로도 잡지 못했다.")
B("**지금 차의 빈틈.** ABS·ESC 는 미끄러지는 바퀴를 잡지만 노면 마찰을 다른 장치에 알리지 않는다. 긴급제동·간격 유지는 앞차 거리로 제동 시점을 정하는데, "
  "마찰을 모르면 마른 노면 기준이라 빙판에서 늦다. 60 km/h 에서 마른 노면은 38 m 앞에서 밟으면 서지만 블랙아이스(μ 0.16)는 105 m 가 필요하다.")
P("그래서 1차는 블랙아이스를 직접 맞히려 하지 않고 '결빙 위험 노면 확률 + 반사도 이상 + 기상 맥락'으로 위험도를 올린다. 그 전략이 실패할 때를 위해 2차가 있다.")
FIGURE(FIG / "real_failure_miss.jpg", 15, "보드가 실제로 놓친 실사진 — 정답은 모두 블랙아이스인데 '정상'으로 판정했다(위험도 0.01~0.04)")

# ═════════════════════════════ 3. 구성 ═════════════════════════════
H1("3. 시스템 구성")
H2("3.1 이중 안전망 구조")
FIGURE(FIG / "architecture.png", 15.5, "시스템 구조 — STM32N6 한 보드에 1차(NPU)와 2차(ThreadX)가 공존하고, 라즈베리파이가 기상 맥락·브리지를 맡는다")
TABLE("두 방어의 역할", ["", "1차 방어 · 예측", "2차 방어 · 반응"],
      [["센서", "전방 카메라", "IMU (횡가속 · yaw rate · 종가속)"],
       ["판단", "노면 4분류 + 반사도 + 기상 맥락 → 위험도", "칼만 필터 + 자전거 모델 잔차 → 미끄러짐 확정, 노면 확인 탭 → 마찰 → 제동 시점"],
       ["시점", "빙판에 닿기 전", "카메라가 애매할 때·미끄러지기 시작한 뒤"],
       ["실행", "STM32N6 NPU · 프레임당 25 ms", "STM32N6 + ThreadX · 최악 22.6 µs"]],
      [2.2, 6.6, 8.2])
P("ThreadX 우선순위는 IMU 융합 스레드가 3, NPU 프레임 스레드가 4다. 25 ms 짜리 추론이 돌고 있어도 IMU 판정이 선점한다(혼합 임계도 구조).")
H2("3.2 하드웨어")
TABLE("하드웨어 구성", ["구분", "정식 명칭", "역할"],
      [["메인 보드", "STM32N6570-DK (MB1939-N6570-C02 + MB1860B)", "1차 NPU 추론(Neural-ART @ 1 GHz) + 2차 ThreadX 실시간 판정·제어"],
       ["카메라", "ST AI Camera module MB1854B", "1차 방어 입력. 보드에 직결되어 NPU 로 바로 들어간다"],
       ["호스트", "Raspberry Pi 5 16 GB + AI HAT+ 2 (Hailo-10H)", "기상 맥락 생성, ZMQ 브리지 (대회 필수 보드)"],
       ["센서", "Intel RealSense D435i", "RGB·뎁스·IMU — 1차 방어 카메라 시연"],
       ["운전 장치", "Moza R3 휠 + SRP 페달", "사람이 직접 CARLA 를 운전하는 HIL"],
       ["렌더", "RTX 5090 · RTX 3090 워크스테이션", "CARLA 0.9.16 물리·센서"],
       ["저장·전원", "NVMe 1 TB SSD · 보조배터리", "데이터·로그 / 콘센트 없이 같은 구성"]],
      [2.4, 6.0, 8.6])

# ═════════════════════════════ 4. 변경점 ═════════════════════════════
H1("4. 개발계획서 대비 변경점")
TABLE("변경점과 이유", ["항목", "계획서(6월)", "실제 구현", "바꾼 이유"],
      [["2차 실행 환경", "Pi 5 + 리눅스", "**STM32N6 + ThreadX**", "리눅스 최악 깨어남 지연 5,790 µs — 20 ms 제어 주기의 29 % 를 지터 한 번이 먹는다"],
       ["1차 NPU", "Hailo-10H 단독", "**Neural-ART(N6) 주**", "Hailo Dataflow Compiler 미확보. 1·2차를 한 보드에 올려 선점 구조를 증명"],
       ["검증", "1/5 차량 저마찰 실측", "**HIL + 실사진 69,358장**", "실물 빙판 재현이 어렵다. 펌웨어는 실물 보드, 비전은 실제 도로 사진으로 측정"],
       ["2차 출력", "미끄러진 뒤 제동·조향", "**+ 마찰 인식 제동 시점**", "실험 결과 결과를 정한 것은 제동 시작 시점이었다 (7장)"]],
      [2.4, 3.0, 3.6, 8.0])

# ═════════════════════════════ 5. 1차 ═════════════════════════════
H1("5. 1차 방어 — 카메라가 미리 본다")
H2("5.1 데이터 전략")
P("데이터셋을 먼저 고르지 않았다. 모델의 각 헤드가 무엇을 정답으로 배워야 하는지를 정의한 뒤 공개 데이터셋을 매핑했다. "
  "주 학습은 RSCD(27클래스 → 정상·젖음·결빙 위험·포트홀 4클래스)이고, 반사도 헤드는 RoadSaW 의 실측 수막 두께, 국내 도로 검증은 AI Hub, "
  "2차 칼만 공분산은 PVS 실측 IMU 를 근거로 했다. 보조로 ROAD·CADC·WADS 를 썼다.")
TABLE("모델 헤드별 정답과 데이터셋", ["방어", "필요 역할", "필요한 정답", "데이터셋", "활용"],
      [["1차", "노면 상태 분류", "건조/젖음/눈/결빙 클래스", "**RSCD**", "주 학습"], ["1차", "반사도·젖음 정도", "실측 수막 두께", "**RoadSaW**", "반사도 헤드 근거"],
       ["1차", "국내 도로·악천후", "눈·비·안개 + 조도", "**AI Hub**", "도메인 검증"], ["2차", "IMU 미끄러짐 감지", "실측 6축 IMU + 노면 라벨", "**PVS**", "칼만 필터 검증"],
       ["2차", "카메라–IMU 융합", "RGB-IMU 동기, 악조건", "**ROAD**", "비교 기준선"], ["보조", "폭설 실주행", "눈 장면·포인트 라벨", "**CADC · WADS**", "악천후 점검"],
       ["보조", "차선 가시성", "자동 라벨", "**CARLA**", "합성 데이터 생성"]], [1.5, 3.4, 4.6, 3.3, 4.2])
P("RSCD 는 약 100만 장(공개 서브셋 37만 장)의 실도로 노면 패치로, 27클래스(마찰 6 × 재질 4 × 요철 3)를 우리 4클래스로 묶었다 — "
  "정상 ← 건조, 젖음 ← 젖음·물 고임, 결빙 위험 ← 결빙·녹은 눈, 포트홀 ← 요철 '심함'. RoadSaW 에는 눈·얼음이 없고 ROAD 의 클래스는 노면 종류라 결빙이 없다 — "
  "이 한계를 알고 역할을 나눠 썼다.")
FIGURE(grid([ROOT / "sw/docs/images/ds_rscd_classes.jpg", ROOT / "sw/docs/images/ds_roadsaw.jpg", ROOT / "sw/docs/images/ds_pvs.jpg"], "datasets", cell_h=300,
            labels=["RSCD", "RoadSaW", "PVS"]), TEXT_W, "주요 데이터셋 — RSCD 클래스별 샘플, RoadSaW 젖음 단계, PVS IMU 장착 (각 데이터셋 논문·공식 자료에서 인용)")
H2("5.2 모델과 보드 배포")
TABLE("RoadNet 과 int8 배포", ["항목", "값"],
      [["백본", "MobileNetV3-Small (ImageNet 사전학습)"], ["입력 / 출력", "1×3×224×224 int8 / 4클래스 로짓 + 반사도 헤드"],
       ["양자화", "int8 PTQ(QDQ), stem 등화 + 클리핑으로 정확도 회복, u8→i8 변환"],
       ["NPU", "Neural-ART @ 1 GHz, 추론 25 ms"], ["카메라 ROI", "전방 7.9 ~ 42.2 m (FOV 60°, 피치 −12°)"]], [4.0, 13.0])
H2("5.3 위험도 융합")
P("risk = α · p_ice + β · 반사도 + γ · (1 − 차선 가시성). 기상·위치 맥락(교량·새벽 등)이 가중치를 바꾼다. "
  "쓸 수 없는 신호는 0 을 넣지 않고 분모에서 빼서 재정규화한다(0 은 중립이 아니라 최솟값이라 상한이 잘린다). "
  "문턱은 초기 0.441 에서 실사진 전량 스윕으로 **0.603** 으로 재교정해 젖은 노면 오경보를 40 % 에서 0.8 % 로 내렸다.")
TABLE("기상·위치 맥락에 따른 융합 가중치", ["맥락", "α (분류)", "β (반사도)", "γ (차선)"],
      [["교량 · 새벽 등 결빙 위험 높음", "0.35", "**0.45**", "0.20"], ["일반", "**0.50**", "0.30", "0.20"], ["저위험", "**0.55**", "0.15", "0.30"]], [7.0, 3.3, 3.4, 3.3])
P("폭우에서는 빙판이 없는 대조군에서도 위험도가 0.96 까지 올라 어떤 문턱으로도 막을 수 없었다. 그래서 강수 게이트로 1차 방어를 스스로 끄고 2차에 맡긴다 — 못 막을 경보를 내는 것보다 '지금은 못 본다'고 말하는 편이 낫다.")
FIGURE(grid([A / "w_clear.jpg", A / "w_rainnight.jpg", A / "w_snow.jpg", A / "w_heavyrain.jpg"], "weather4", cell_h=280,
            labels=["맑은 낮", "비 오는 밤", "눈", "폭우"]), TEXT_W, "맥락 계층이 보는 입력 — 같은 빙판·같은 주행선에서 날씨만 바꿨다")
H2("5.4 실사진 대규모 평가")
TABLE("실제 도로 사진 69,358장 — STM32N6 보드 판정", ["실제 노면", "장수", "정답률", "경보율(문턱 0.603)"],
      [["블랙아이스", "6,340", "96.1 %", "**96.6 %**"], ["마른 노면", "19,018", "80.7 %", "**0.3 %**"],
       ["젖은 노면", "34,440", "84.6 %", "**0.8 %**"], ["포트홀", "9,560", "90.2 %", "0.3 %"]], [4.0, 3.0, 3.0, 7.0])
FIGURE(FIG / "real_threshold.jpg", 12.5, "운영 문턱 결정 — 69,358장 전량 스윕에서 절벽 구간을 찾아 0.603 으로 정했다")
FIGURE(FIG / "real_confusion.jpg", 9.5, "혼동 행렬 (행 = 실제 노면, 열 = 보드 판정)")
H2("5.5 CARLA 동작")
P("CARLA Town04, 맑은 낮, 40 km/h, 마찰 0.02. 6.10 s 에 위험도 0.443 으로 1차 경보, 7.14 s 에 빙판 23.6 m 앞에서 정지했다. "
  "같은 주행선에서 날씨 11종을 모두 돌려 인식·정지 14회, 경고 늦음 7회, 미인식 10회, 오경보 1회였다.")
FIGURE(filmstrip(V / "v_primary.mp4", "primary"), TEXT_W, "1차 방어 동작 — 각 칸 왼쪽 전방 카메라(NPU 입력), 오른쪽 조감. 접근 → 경보 → 빙판 앞 정지")
FIGURE(FIG / "weather_1차_카메라경고.jpg", 15.5, "날씨 11종에서 1차 경보가 난 같은 순간 — 같은 주행선·같은 빙판에서 날씨만 바꿨다")

# ═════════════════════════════ 6. 2차 ═════════════════════════════
H1("6. 2차 방어 — 카메라가 놓치면 차량 거동으로 받는다")
H2("6.1 미끄러짐 판정")
B("IMU 50 Hz 에서 횡가속 a_y · yaw rate · 종가속을 받는다.")
B("2상태 칼만 필터(값 · 변화율)로 잡음을 거른다.")
B("자전거 모델로 '이 속도·이 조향이면 나와야 할 yaw rate'를 계산한다(고속 언더스티어 보정, 조향 지연 τ 0.06 s).")
B("실측과 모델의 잔차가 **타원 밖**이면 미끄러짐 — 3샘플(60 ms) 연속이면 확정한다. 직사각형 규칙은 모서리에서 저속 지연이 있어 타원으로 바꿨다.")
TABLE("미끄러짐 판정 파라미터", ["항목", "값", "이유"],
      [["축거 L", "2.7 m", "자전거 모델"], ["조향 지연 τ", "**0.06 s**", "급조향 때 '모델이 즉답한다'는 가정이 오탐을 만들었다"],
       ["횡가속 · yaw 오차 임계", "0.30 g · 0.35 rad/s", "타원의 두 반축"], ["판정 규칙", "**타원**, 3샘플(60 ms) 연속", "직사각형은 모서리에서 저속 지연이 생겼다"],
       ["저마찰 트리거", "제동 ≥ 0.3 이 0.3 s 지속 + 감속 < 1.2 m/s²", "필터 지연 구간의 오탐 방지"]], [4.0, 5.0, 8.0])
FIGURE(FIG / "rule_boundary.jpg", 12, "미끄러짐 판정 경계 — 직사각형 규칙의 모서리 지연을 없애려 타원 규칙으로 바꿨다")
P("CARLA Town04, 맑은 낮, 40 km/h, 마찰 0.08, 1차 방어를 끈 상태(미인식 가정)로 38.7 km/h 진입했다. 7.62 s 에 STM32N6 RTOS 가 미끄러짐을 확정했고(진입 1.78 s 만, a_y −0.364 g, yaw 오차 −0.282), "
  "hard_stop 뒤 왼쪽 차로가 비어 evade_left 로 전환해 11.68 s 에 차선 안에서 정지했다. 전체 2차 방어 발동 62회(보드 판정 53 · 호스트 9), 진입→확정 평균 2.41 s, 확정→정지 평균 3.78 s.")
FIGURE(filmstrip(V / "v_secondary.mp4", "secondary"), TEXT_W, "2차 방어 동작 — 카메라가 놓친 빙판 진입 → STM32N6 이 미끄러짐 확정 → 차선 안 정지 (각 칸 왼쪽 전방 카메라, 오른쪽 조감)")
P("C 코어(sw/fw/npu_lib/slip_core.h)는 HAL·OS 비의존이라 호스트에서도 컴파일되고, 파이썬 참조 구현과 같은 판정을 내리는지 매 커밋 검증한다.")
H2("6.2 왜 리눅스가 아니라 RTOS 인가")
TABLE("같은 연산의 지연", ["플랫폼", "표본", "중앙값", "최악"],
      [["Pi 5 + Linux · 유휴", "20,000", "69 µs", "**5,790 µs**"], ["Pi 5 + Linux · 부하", "20,000", "68 µs", "5,429 µs"],
       ["STM32N6 + ThreadX (응답 전체)", "49,405", "12.4 µs", "**22.6 µs**"]], [7.0, 3.0, 3.0, 4.0])
FIGURE(FIG / "latency_cdf.jpg", 12, "지연 분포(CDF) — 리눅스는 꼬리가 길고, ThreadX 는 최악 22.6 µs 안에 끝난다")
FIGURE(FIG / "schedule_rtos.jpg", 13, "RM 스케줄 — IMU 융합(우선순위 3)이 NPU 프레임(4)을 선점한다. 뒤집으면 25 ms 추론이 20 ms 주기를 막는다")
P("리눅스는 평균이 빠르지만 꼬리가 길다. 안전 기능의 기준은 평균이 아니라 최악이다. 우선순위를 뒤집으면 25 ms 추론이 20 ms 주기를 막아 스케줄 자체가 불가능해진다.")
H2("6.3 방어가 없으면 / 2차 방어가 있으면")
P("1차·2차를 모두 끈 기준선 9건은 전부 차선을 벗어나 88~99° 가로로 미끄러졌고, 5건은 42~52 km/h 로 방호벽에 부딪혔다(충돌 직전 속도). "
  "같은 조건에서 STM32N6 RTOS 2차 방어는 미끄러짐을 확정하고 차선 안에서 세웠다. 2차 방어 발동 62회, 진입→확정 평균 2.41 s, 확정→정지 평균 3.78 s.")
FIGURE(filmstrip(V / "v_nodefense.mp4", "nodefense"), 14.5, "방어 없음 기준선 — 차선 이탈 → 가로 미끄러짐 → 42 km/h 로 방호벽 충돌 (화면 글씨 12 km/h 는 충격 뒤 속도)")
FIGURE(filmstrip(V / "v_compare.mp4", "compare"), TEXT_W, "같은 조건 직접 비교 — 각 칸 왼쪽 방어 없음, 오른쪽 STM32N6 RTOS 2차 방어. 같은 날씨·주행선·빙판·주변 차량 8대")
FIGURE(grid([A / "c_clearnoon.jpg", A / "c_clearnight.jpg", A / "c_wetnoon.jpg", A / "c_snow.jpg"], "compare4", cols=2, cell_h=260,
            labels=["맑은 낮", "맑은 밤", "젖은 노면", "눈"]), 15, "날씨 4종에서 같은 비교 — 각 쌍의 마지막 장면 (왼쪽 방어 없음, 오른쪽 RTOS 2차 방어)")

# ═════════════════════════════ 7. 확장 ═════════════════════════════
H1("7. 2차 방어 확장 — 노면 마찰을 재서 제동 시점을 고친다")
H2("7.1 '어떻게 밟나'가 아니라 '언제 밟나'")
P("먼저 미끄러진 뒤의 제어를 다듬는 실험을 했다. 곡선 빙판에서는 마찰별 제어표 후보 78개가 고정 게인과 충돌 직전 속도 중앙값 0.2 km/h 안으로 같았고, "
  "직선 빙판 위 정지 차량 장면에서는 2차 방어를 꺼도 결과가 같았다. 빙판에서는 누가 밟든 감속이 마찰 한계(μg)에서 막히고, 펄스 제동은 실차 ABS 가 이미 하는 일이다. "
  "결과를 정한 것은 **제동을 시작한 시점**이었다(60 km/h 에서 9 s 전에 밟으면 전부 충돌, 13 s 전이면 정지).")
FIGURE(FIG / "gs_curve_impact.png", 15.5, "곡선 빙판 — 방어 없음 · 고정 게인 · 무작위 제어표 후보의 충돌 직전 속도. 후보는 고정 게인을 넘지 못했다")
FIGURE(FIG / "si_noevade.png", 15.5, "직선 빙판 위 정지 차량 — 판단 방식을 바꿔도 결과가 같다. 운전자 제동 시점(아래 숫자)이 결과를 정했다")
TABLE("기존 기술과의 차이", ["", "하는 일", "못 하는 것"],
      [["ABS · ESC", "미끄러지는 바퀴를 잡고 차체 회전을 막는다", "마찰을 다른 장치에 알리지 않는다"],
       ["긴급제동 · 간격 유지", "앞차 거리로 제동 시점을 정한다", "마찰을 모르면 마른 노면 기준이라 빙판에서 늦다"],
       ["노면 마찰 추정(일부 양산차)", "바퀴 신호로 미끄러움을 추정해 경고·공유", "제동 시점까지 바로 잇는 경우는 드물다"],
       ["**IcePredict 2차 방어**", "**마찰을 직접 재서 제동 시작 시점을 노면에 맞춘다**", "탭을 운전자가 느낀다"]], [4.0, 6.5, 6.5])
H2("7.2 노면 확인 탭과 마찰 인식 제동 시점")
P("빙판은 살짝만 밟아도 감속이 거기서 막힌다. 그 막힌 값이 곧 노면 마찰이다. 무거운 계산(마찰별 실측 감속 표)은 CARLA 에서 미리 끝내고, 보드는 20 ms 마다 표만 본다.")
B("**언제 확인하나** — ① 1차 카메라 위험도가 회색 구간(0.3~0.6)에 0.2 s 이상이면 카메라가 본 지점(19 m 앞)에 도착했을 때, ② 영하 근처(3 °C 이하) 결빙 취약 구간 안이면 2.5 s 마다 최대 3번. 안전 조건: 30 km/h 이상, 직진, 운전자 제동 없음, 앞차 여유.")
B("**노면 확인 탭** — 가속 페달을 0.2 s 떼고 브레이크 0.3 을 0.4 s. 마지막 0.25 s 평균 감속을 표로 되돌려 μ 를 잰다. 마른 노면이면 60 ms 만에 끝낸다(속도 손실 13 → 5 km/h).")
B("**20 ms 마다 판단** — 정지 거리 = 0.3 s × v + v² / (2 × 0.9 × 실측 감속) + 4 m 를 앞 물체 거리와 비교해 순항·가속 금지·예비 제동·전력 제동·정지 유지 중 하나를 고른다.")
FIGURE(FIG / "env_table.png", 14, "제동을 시작해야 하는 거리 — CARLA 실측 감속 기준. 60 km/h 마른 노면 38 m, 블랙아이스(μ 0.16) 105 m")
FIGURE(FIG / "env_probe_timeline.png", 15, "대표 주행 — 결빙 취약 구간 첫 탭은 마른 노면 '이상 없음', 다음 탭이 빙판 μ 0.15. 1차 카메라 위험도는 내내 낮았다")
H2("7.3 CARLA 비교 결과")
P("간격 유지 크루즈(ACC) 주행 중 빙판 위 정지 차량 장면에서, 같은 C 코드에 마찰 값만 다르게 넣었다. 비교 대상은 마른 노면 μ 를 가정하고, 2차 방어는 탭으로 잰 μ 를 쓴다.")
TABLE("마른 노면 기준 vs 마찰 인식 (마찰 3종 × 속도 3종)", ["", "마른 노면 기준", "마찰 인식"],
      [["정지", "2 / 9", "**8 / 9**"], ["충돌", "7 / 9 (8 ~ 61 km/h)", "1 / 9 (48 km/h, 필요 거리 295 m > 150 m)"],
       ["정지 위치", "—", "7회 4.1 ~ 4.2 m 앞 (목표 여유 4 m), 1회 0.8 m 앞"]], [3.5, 5.5, 8.0])
FIGURE(FIG / "env_sweep.png", 15.5, "충돌 직전 속도(0 = 정지) — 마찰 3종 × 속도 3종")
FIGURE(filmstrip(V / "v_envelope.mp4", "envelope"), 12.5, "비교 장면 — 각 칸 위: 마른 노면 기준(30 km/h 충돌), 아래: 탭으로 마찰 측정(4.1 m 앞 정지)")
TABLE("탭 판단 시험 (1차 모델이 실제로 판정)", ["조건", "결과"],
      [["빙판 + 결빙 취약 구간 · 맑은 낮·밤·젖은 날씨", "3 / 3 정지 — 첫 탭 마른 노면, 다음 탭에서 빙판"],
       ["빙판 · 구간 정보 없음", "카메라가 놓쳐 탭 없음 → 30 km/h 충돌"],
       ["빙판 없음 · 구간 안 · 영하", "탭 3번 모두 '이상 없음' — 잘못된 제동 0"], ["빙판 없음 · 영상 8 °C", "탭 없음"]], [8.0, 9.0])
P("1차 카메라 모델은 이 CARLA 장면들에서 빙판을 거의 보지 못했다(위험도 최대 0.05). 그래서 카메라 계기만으로는 탭이 나가지 않았고, 지도의 결빙 취약 구간 정보가 그 공백을 메운다.")
H2("7.4 기초 작업 — CARLA 마찰 보정과 보드 마찰 추정기")
P("마찰 인식 제동의 표는 CARLA 실측에서 나왔다. friction 9종 × 시험 9종으로 바퀴가 실제로 낸 최대 가속을 재서, friction 0.3 이하에서 실효 μ ≈ 0.81 × friction 임을 확인했다(0.5 부터는 제동 토크가 먼저 한계). "
  "보드 C 코어에는 미끄러짐 확정 뒤 펄스 제동 중의 합가속으로 μ 를 추정하는 기능(SLIP_MU_EST)을 넣었고, 실효 μ 0.41 이하에서 펄스 3회(약 1초) 안에 ±30 % 로 수렴했다.")
FIGURE(grid([FIG / "gs_mu_calib.png", FIG / "gs_mu_converge.png"], "mu_basis", cell_h=420), TEXT_W, "왼쪽: CARLA friction 과 실효 마찰(최대 가속/g). 오른쪽: 보드 마찰 추정기의 수렴 — 펄스 제동 로그 재생")
P("이 과정에서 CARLA 결함 두 가지를 찾았다. 차량에 준 타이어 마찰 설정은 효과가 없어 마른 노면은 사실 기본값(횡 약 2 g)이었고, 빙판 트리거를 벗어나면 다시 그 값으로 돌아가 옆으로 미끄러지던 차가 2.5 g 로 급정지했다. "
  "그래서 검증 주행은 빙판을 길고 넓게 깔고, 비교는 빙판 안 구간으로만 잘랐다.")
FIGURE(FIG / "gs_friction_trigger.png", 12, "CARLA 마찰 트리거 결함 — 트리거 밖으로 나가면 마찰이 기본값으로 돌아간다")
P("**측정 함정.** CARLA 는 마찰 트리거 안에 다른 차량이 있으면 자차가 트리거 밖에 있어도 처음부터 그 마찰을 건다. 정지 차량을 트리거 밖에 두고 모든 실험을 다시 했고, 주행마다 출발 가속으로 마른 출발을 확인했다.", size=9, color=GREY)

# ═════════════════════════════ 8. 안전 전용 RTOS ═════════════════════════════
H1("8. 안전 전용 RTOS — 메인 컴퓨터가 멈춰도 2차 방어는 산다")
P("STM32 의 몫은 '더 좋은 제어'가 아니라 **고장 격리**다. 곡선 빙판에서 메인 컴퓨터(파이)를 멈추자, 2차 방어가 파이에 있을 때는 방어 없음과 같은 53.1 km/h 로, "
  "안전 전용 RTOS 에 있을 때는 고장 없음과 같은 39.4 km/h 로 충돌했다. 파이 심장박동이 0.12 s 끊기면 축소 모드로 들어가 카메라·지도 입력을 버리고 이미 잰 μ 로 판단을 계속한다.")
FIGURE(FIG / "si_fault.png", 15.5, "곡선 빙판에서 파이를 죽였을 때 충돌 직전 속도 — 2차 방어 위치에 따라 결과가 갈린다")
P("STM32N6 보드가 없어 확장 기능은 같은 C 코드를 **라즈베리파이 격리 코어**(코어 3 · SCHED_FIFO · mlockall, 안전 전용 RTOS 대체)에서 돌려 측정했다. "
  "리눅스 커널 자체가 멈추면 같이 멈춘다 — 그래서 실차 설계는 별도 MCU 다.")
TABLE("전 코어 부하에서 20 ms × 6,000회", ["구성", "평균 지연", "최악 지연", "마감 놓침"],
      [["일반 리눅스 프로그램", "2.6 ms", "206 ms", "211"], ["실시간 우선순위 + 코어 고정", "0.04 ms", "0.23 ms", "0"],
       ["**파이 격리 코어 (RTOS 대체)**", "0.018 ms", "**0.06 ms**", "**0**"]], [7.0, 3.0, 3.5, 3.5])
FIGURE(filmstrip(V / "v_island.mp4", "island"), TEXT_W, "빙판 확인 직후 파이 메인 프로그램을 실제로 강제 종료(SIGKILL) — 격리 코어의 2차 방어가 혼자 4.1 m 앞에 세운다")

# ═════════════════════════════ 9. 시연 ═════════════════════════════
H1("9. 시연 구성")
TABLE("시연 세 장면", ["장면", "장비", "보여주는 것"],
      [["① 1차 방어", "RealSense D435i", "빙판 사진을 비추면 1차 모델이 실시간으로 '블랙아이스 감지 — 감속하세요'"],
       ["② 2차 방어", "Moza 휠·페달 + CARLA", "2차 방어 OFF 면 평소처럼 밟아도 미끄러져 충돌, ON 이면 탭 → 경고 → 필요한 순간 자동 제동"],
       ["③ 고장 격리", "라즈베리파이 5", "주행 중 파이 메인 프로그램을 실제로 죽여도 격리 코어의 2차 방어가 계속 판단"]], [2.8, 4.2, 10.0])
P("사람이 운전해도 시뮬레이션은 실시간의 0.996 배로 돈다. 2차 방어 판단은 파이 격리 코어가 20 ms 마다 내고, 캠퍼스 와이파이 왕복 지연은 중앙 4.8 ms 다. "
  "Moza 휠은 중앙 복귀 스프링, 탭 때 짧은 충격, 빙판 확정 때 진동으로 실차에서 몸으로 느끼는 감속을 대신한다.")
FIGURE(grid([A / "demo_primary.jpg", A / "demo_island.jpg"], "demo2", cell_h=420), TEXT_W,
       "왼쪽: 1차 방어 시연 화면 — 카메라 입력과 노면 분류 확률·위험도·경고. 오른쪽: Moza 시연 화면 — 빙판 감지 배너, 2차 방어 패널, 파이 메인 프로그램 정지 표시")
P("시연 영상(1분 50초)은 네 편으로 찍었다. ① RealSense 로 실제 노면 사진을 비춰 마른·젖은 노면과 블랙아이스를 판정, "
  "② 자동 주행 중 RealSense 가 빙판 사진을 보자 1차 방어가 빙판 **7.9 m 앞에 정지**, ③·④ 같은 장면을 2차 방어 OFF/ON 으로 — "
  "OFF 는 브레이크 지점(앞차 65 m)에서 밟고 핸들을 꺾어도 미끄러져 **앞차와 20 km/h 충돌**, ON 은 빙판 시작점의 탭으로 μ 0.21 을 재고 미리 제동해 **앞차 28 m 앞 정지**. "
  "③·④ 의 차는 자동 운전자가 몰고, 화면 오른쪽에 같은 순간 Moza 를 조작하는 운전자 영상을 붙였다. CARLA 차량에 없는 ABS·ESC 는 대역으로 넣어 OFF/ON 에 똑같이 적용했다.")
FIGURE(grid([A / "demo_real_rs.jpg", A / "demo_real_moza.jpg", A / "demo_real_pedal.jpg"], "demo_real", cell_h=420,
            labels=["RealSense 노면 판정", "Moza 조작", "휠·페달"]), TEXT_W,
       "실제 시연 — RealSense D435i 로 빙판 사진을 비추면 '블랙아이스 감지 — 감속하세요'(왼쪽), Moza 휠·페달로 CARLA 를 조작하는 운전자(가운데·오른쪽)")
FIGURE(grid([A / "demo_res_1st.jpg", A / "demo_res_off.jpg", A / "demo_res_on.jpg"], "demo_res", cell_h=300,
            labels=["1차: 빙판 앞 정지", "2차 OFF: 충돌", "2차 ON: 28 m 앞 정지"]), TEXT_W,
       "시연 영상의 결과 장면 — 오른쪽 패널에 1차 카메라(RealSense) 화면 또는 같은 순간의 Moza 조작 영상이 들어간다")

# ═════════════════════════════ 10. 재현 ═════════════════════════════
H1("10. 검증 방법과 재현성")
B("실물 STM32N6 보드가 실제 펌웨어를 돌리고 CARLA 가 센서·물리를 제공하는 HIL 로 주행 108회를 측정했다. 주행별 이벤트 JSON 을 저장소에 함께 올렸다.")
B("1차 성능은 시뮬레이션 화면이 아니라 실제 도로 사진 69,358장을 보드에 넣어 쟀고, 장당 판정을 남겨 문턱을 바꿔 재계산할 수 있다.")
B("2차 방어 C 코드 = 파이썬 참조 = 파이 격리 코어 출력이 샘플 단위로 같다. 확장 기능은 컴파일 플래그(기본 0) 뒤에 있고, 끄면 작업 전과 출력이 비트 단위로 같다.")
B("모든 수치는 저장소의 스크립트로 다시 만들 수 있다. 근거 문서 23건, 결과 그림 41건.")
P("주행마다 무엇을 보고 판단했는지 남긴다. 시맨틱 라이다 64채널에서 빙판 구간과 차량 3D 박스를, 카메라 영상에서 도로·차선·빙판·차량 2D 분할 라벨을 자동으로 만든다(AI Hub 형식 참고).")
FIGURE(FIG / "triplet_rscdtex_visual_00320.jpg", TEXT_W, "카메라 · 2D 분할 라벨 · 라이다 — 도로/차선/빙판/차량 라벨 자동 생성")

# ═════════════════════════════ 11. 한계 ═════════════════════════════
H1("11. 한계와 향후 계획")
TABLE("한계", ["항목", "내용"],
      [["경보 거리", "카메라 ROI 7.9 ~ 42.2 m. 모델 최소 80 px 과 ROI 해상도 91 px 이 경보 거리를 정한다 — 해상도를 올리는 것이 먼저다"],
       ["2차는 예방이 아니다", "60 km/h 에서 1차가 경고했지만 제동이 늦어 진입했고, 2차가 0.26 s 만에 개입했는데도 39.2 km/h 로 앞차와 충돌했다"],
       ["오경보·도메인 갭", "밤에 41 m 밖 빙판을 두고 오경보가 났고, RSCD 질감을 입힌 CARLA 노면에는 같은 모델이 반응하지 않는다"],
       ["탭", "마찰 측정 탭을 운전자가 느낀다(마른 노면 5 km/h 손실)"],
       ["시뮬레이션·측정", "확장 결과는 CARLA 시뮬레이션이고, 실시간 측정은 파이 격리 코어 대체다. CARLA 타이어는 잠겨도 마찰이 줄지 않는다"]], [3.6, 13.4])
FIGURE(grid([A / "falsealarm_c.jpg", A / "domaingap_c.jpg"], "limits", cell_h=360, labels=["밤 오경보", "질감 도메인 갭"]), 15,
       "한계 사례 — 왼쪽: 41 m 밖 빙판을 두고 어두운 노면을 얼음으로 본 오경보(위험도 0.52). 오른쪽: RSCD 질감을 입힌 CARLA 노면에 모델이 반응하지 않았다")
P("**향후 계획**", after=2)
B("경보 거리 확대 — ROI 해상도를 올려 1차가 보는 거리를 늘린다.")
B("보드 실측 — 마찰 인식 제동 시점·파이 감시를 STM32N6 ThreadX 에서 다시 잰다.")
B("운전자가 못 느끼는 마찰 측정 — 앞축 구동·뒷축 제동으로 차 전체 가속을 바꾸지 않고 마찰만 잰다.")
B("놓친 장면 재학습 — 2차 방어가 확정한 빙판 직전 카메라 영상을 1차 방어 학습 데이터로 쓴다.")

# ═════════════════════════════ 부록 ═════════════════════════════
H1("부록. 근거 문서 목록")
docs_ = sorted(p.name for p in (ROOT / "sw/docs/evidence").glob("[0-9][0-9]_*.md"))
TABLE("근거 문서 (저장소 sw/docs/evidence)", ["번호", "문서"], [[d[:2], d[3:-3].replace("_", " ")] for d in docs_], [2.0, 15.0], size=8.5)
P(f"소스 코드·근거 문서·영상: {REPO}", size=9, color=BLUE)

doc.save(OUT)
print(f"저장: {OUT} ({OUT.stat().st_size / 1e6:.1f} MB) — 그림 {_fig_n[0]}, 표 {_tab_n[0]}")
