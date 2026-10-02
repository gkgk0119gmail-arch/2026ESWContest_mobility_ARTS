#!/usr/bin/env python3
"""PDF 변환 전처리: 임베드 영상을 포스터 그림으로 바꾼 사본을 만든다.

LibreOffice 는 영상을 PDF 안에 그대로 묻어 27 MB 짜리 PDF 가 나온다(심사용 PDF 에 영상은 재생되지도 않는다).
영상 pic 에서 a:videoFile · p14:media · 재생 하이퍼링크 · p:timing 만 떼어내면 blipFill 의 포스터가 남아
보통 그림이 된다. 원본 pptx 는 건드리지 않는다.

사용: python3 pdf_prep.py 원본.pptx 출력.pptx
"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
n = 0
for slide in prs.slides:
    for sh in list(slide.shapes):
        el = sh._element
        nvPr = el.find(f"{qn('p:nvPicPr')}/{qn('p:nvPr')}")
        if nvPr is None or nvPr.find(qn("a:videoFile")) is None:
            continue
        for tag in ("a:videoFile", "p:extLst"):
            for c in nvPr.findall(qn(tag)):
                nvPr.remove(c)
        cNvPr = el.find(f"{qn('p:nvPicPr')}/{qn('p:cNvPr')}")
        for c in cNvPr.findall(qn("a:hlinkClick")):
            cNvPr.remove(c)
        n += 1
    for t in slide._element.findall(qn("p:timing")):
        slide._element.remove(t)
prs.save(dst)
print(f"영상 {n}개를 포스터 그림으로 바꿔 저장: {dst}")
