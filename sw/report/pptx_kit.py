#!/usr/bin/env python3
"""build_pptx.py 가 쓰는 디자인 시스템 + 레이아웃 도구.

디자인 기준은 현대자동차 브랜드 스타일가이드(Hyundai European Website Styleguide, 2020)다.
  · 색: Hyundai Blue #002C5F 를 큰 면(표지·강조 띠)에, Active Blue #00AAD2 / Active Red #E63312 를
    1차(예측·AI) / 2차(반응·RTOS) 의 의미색으로, Light Sand #E4DCD3 계열을 패널 바탕으로 쓴다.
  · 글꼴: Hyundai Sans 는 사내 전용이라 같은 계열(기하학적 고딕)의 공개 글꼴 Pretendard 를 쓴다.
  · 모양: 둥근 모서리·테두리·색 띠 없이, 사진은 모서리 그대로 격자에 붙이고 여백으로 구분한다.

좌표는 전부 px(1280×720) 다. 1280×720 px = 13.333in × 7.5in (96 dpi) 이므로 px/96 = 인치.
"""
from __future__ import annotations

import io
import pathlib
import re

from PIL import Image, ImageOps
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PER_PX = 914400 // 96          # 9525

# ── 현대 브랜드 색 ───────────────────────────────────────────────────────────
HBLUE  = RGBColor(0x00, 0x2C, 0x5F)   # Hyundai Blue   — 주색
ABLUE  = RGBColor(0x00, 0xAA, 0xD2)   # Active Blue    — 1차 방어(예측·AI)
SKY    = RGBColor(0xAA, 0xCA, 0xE6)   # Sky Blue       — 파란 면 위의 보조 글자
ARED   = RGBColor(0xE6, 0x33, 0x12)   # Active Red     — 2차 방어(반응·RTOS)·위험
SAND   = RGBColor(0xA3, 0x6B, 0x4F)   # Hyundai Sand
LSAND  = RGBColor(0xE4, 0xDC, 0xD3)   # Light Sand
PANEL  = RGBColor(0xF6, 0xF3, 0xF2)   # 패널 바탕 (Light Sand 를 더 밝게)
INK    = RGBColor(0x1C, 0x1B, 0x1B)   # Grey 1000 — 제목·본문
BODY   = RGBColor(0x4D, 0x4D, 0x4D)   # Grey 800  — 설명
MUTED  = RGBColor(0x76, 0x76, 0x76)   # Grey 600  — 캡션
FAINT  = RGBColor(0xA6, 0xA6, 0xA6)   # Grey 400  — 쪽수·바닥글
LINE   = RGBColor(0xE0, 0xDC, 0xD8)   # 가는 선
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x1E, 0x8A, 0x4C)   # 성공 표시에만

# 이전 판과의 호환 이름
BLUE, RED, NAVY, DARK, SLATE, ZEBRA = ABLUE, ARED, HBLUE, HBLUE, BODY, PANEL

FONT   = "Pretendard"                 # 본문·굵은 글씨
FONT_X = "Pretendard ExtraBold"       # 제목·큰 숫자
FONT_S = "Pretendard SemiBold"        # 소제목
MONO   = "Consolas"

_COLORS = {"r": ARED, "g": GREEN, "a": ABLUE, "h": HBLUE, "m": MUTED, "w": WHITE, "k": INK, "s": SKY}


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


# ── 인라인 표시 ──────────────────────────────────────────────────────────────
# <b>굵게</b> <br> 줄바꿈 <r>빨강</r> <a>Active Blue</a> <h>Hyundai Blue</h> <g>초록</g> <m>회색</m> <w>흰색</w> <s>Sky</s>
# <rb> <ab> <hb> <gb> 는 굵은 색글씨, <mono> 는 고정폭, <x> 는 ExtraBold.
_TAG = re.compile(r"<(/?)(b|br|x|mono|r|g|a|h|m|w|k|s|rb|gb|ab|hb|mb|wb)>")


def _runs(text: str):
    """(글자, 굵게, 색, 글꼴) 목록. 줄바꿈은 '\\n' 글자로 돌려준다."""
    out, bold, color, face, i = [], False, None, None, 0
    for m in _TAG.finditer(text):
        if m.start() > i:
            out.append((text[i:m.start()], bold, color, face))
        close, tag = m.group(1), m.group(2)
        if tag == "br":
            out.append(("\n", bold, color, face))
        elif tag == "b":
            bold = not close
        elif tag == "x":
            face = None if close else FONT_X
        elif tag == "mono":
            face = None if close else MONO
        elif len(tag) == 1:
            color = None if close else _COLORS[tag]
        else:                                  # rb gb ab hb mb wb
            bold = not close
            color = None if close else _COLORS[tag[0]]
        i = m.end()
    if i < len(text):
        out.append((text[i:], bold, color, face))
    return out


def _set_face(run, face):
    """latin · ea(한글) 글꼴을 같이 지정한다. python-pptx 의 font.name 은 latin 만 바꾼다."""
    run.font.name = face
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def _fill_paragraph(p, raw, size, color, bold, face, spacing_pt=0):
    for frag, b, c, f in _runs(raw):
        for li, chunk in enumerate(frag.split("\n")):
            if li:
                p.add_line_break()
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            r.font.size = Pt(size)
            r.font.bold = bold or b
            r.font.color.rgb = c or color
            _set_face(r, f or face)
            if spacing_pt:
                r._r.get_or_add_rPr().set("spc", str(int(spacing_pt * 100)))


def text(sh_tree, x, y, w, h, body, size=11, color=INK, bold=False, face=None,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.4, space_after=0, tracking=0, wrap=True):
    """글상자 하나. body 는 문단 목록(list[str]) 이거나 문자열. tracking 은 자간(pt)."""
    tb = sh_tree.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = body if isinstance(body, (list, tuple)) else [body]
    for pi, raw in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        _fill_paragraph(p, raw, size, color, bold, face or FONT, tracking)
    return tb


def rect(sh_tree, x, y, w, h, fill=None, line=None, lw=1.0):
    s = sh_tree.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    st = s._element.find(qn("p:style"))      # 테마 참조(그림자 포함)를 떼어낸다 — 선·채움은 위에서 직접 정했다
    if st is not None:
        s._element.remove(st)
    s.text_frame.text = ""
    return s


def hline(sh_tree, x, y, w, color=LINE, lw=0.75):
    """가로 가는 선."""
    ln = sh_tree.add_connector(1, px(x), px(y), px(x + w), px(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    return ln


def veil(sh_tree, x, y, w, h, color=HBLUE, opacity=0.75):
    """반투명 덮개. opacity 1.0 이 완전 불투명."""
    s = rect(sh_tree, x, y, w, h, fill=color)
    srgb = s.fill.fore_color._xFill.find(qn("a:srgbClr"))
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(round(opacity * 100000)))})
    srgb.append(a)
    return s


def label(sh_tree, x, y, w, txt, color=HBLUE, size=9.5):
    """작은 머리표(overline). 굵게 + 자간."""
    return text(sh_tree, x, y, w, size * 1.8, txt, size=size, color=color, bold=True, tracking=0.8, spacing=1.0)


def block(sh_tree, x, y, w, h, title, paras, color=HBLUE, fill=None, size=10.5,
          title_size=11.5, pad=0, body_color=BODY, spacing=1.45, space_after=4):
    """소제목 + 본문. fill 을 주면 PANEL 색 바탕(테두리·띠 없음)."""
    if fill is not None:
        rect(sh_tree, x, y, w, h, fill=fill)
        pad = pad or 16
    tx, ty, tw = x + pad, y + pad, w - 2 * pad
    if title:
        text(sh_tree, tx, ty, tw, title_size * 1.9, title, size=title_size, bold=True, color=color, spacing=1.0)
        ty += title_size * 1.9 + 4
    text(sh_tree, tx, ty, tw, y + h - pad - ty, paras, size=size, color=body_color,
         spacing=spacing, space_after=space_after)


def stat(sh_tree, x, y, w, h, number, caption, color=HBLUE, on_dark=False, num_size=34, cap_size=9.5,
         align=PP_ALIGN.LEFT, fill=None):
    """큰 숫자 + 작은 설명. on_dark 면 파란 면 위에 흰 글씨."""
    if fill is not None:
        rect(sh_tree, x, y, w, h, fill=fill)
    nc = WHITE if on_dark else color
    cc = SKY if on_dark else MUTED
    pad = 14 if fill is not None else 0
    text(sh_tree, x + pad, y + pad, w - 2 * pad, num_size * 1.35, number, size=num_size, color=nc,
         face=FONT_X, align=align, spacing=1.0)
    text(sh_tree, x + pad, y + pad + num_size * 1.35 + 2, w - 2 * pad, h - pad - num_size * 1.35 - 2, caption,
         size=cap_size, color=cc, align=align, spacing=1.3)


# ── 표 ───────────────────────────────────────────────────────────────────────
_NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid


def _cell_border(cell, side, color=None, w_pt=0.75):
    """side: 'L','R','T','B'. color None 이면 선 없음."""
    tcPr = cell._tc.get_or_add_tcPr()
    tag = qn(f"a:ln{side}")
    for old in tcPr.findall(tag):
        tcPr.remove(old)
    ln = tcPr.makeelement(tag, {"w": str(int(w_pt * 12700)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    if color is None:
        ln.append(ln.makeelement(qn("a:noFill"), {}))
    else:
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": str(color)}))
        ln.append(sf)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
    # 순서: lnL lnR lnT lnB 가 fill 보다 앞에 와야 한다
    fills = [c for c in tcPr if c.tag in (qn("a:solidFill"), qn("a:noFill"))]
    for f in fills:
        tcPr.remove(f)
    tcPr.append(ln)
    order = {qn("a:lnL"): 0, qn("a:lnR"): 1, qn("a:lnT"): 2, qn("a:lnB"): 3}
    lns = sorted([c for c in tcPr if c.tag in order], key=lambda c: order[c.tag])
    for c in lns:
        tcPr.remove(c)
    for i, c in enumerate(lns):
        tcPr.insert(i, c)
    for f in fills:
        tcPr.append(f)


def table(sh_tree, x, y, w, head, rows, colw, size=10, row_h=24, head_h=24,
          head_color=HBLUE, zebra=False, first_col_bold=True, head_fill=None, line=LINE, cell_pad=8,
          head_colors=None, first_col_color=None):
    """현대 스타일 표: 세로줄 없음, 머리글은 파란 굵은 글씨 + 아래 굵은 선, 행 사이는 가는 선."""
    n_c, n_r = len(head), len(rows) + 1
    g = sh_tree.add_table(n_r, n_c, px(x), px(y), px(w), px(head_h + row_h * len(rows)))
    tbl = g.table
    tbl.first_row = tbl.horz_banding = tbl.vert_banding = tbl.first_col = tbl.last_row = tbl.last_col = False
    sid = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tbl._tbl.tblPr.makeelement(qn("a:tableStyleId"), {})
        tbl._tbl.tblPr.append(sid)
    sid.text = _NO_STYLE
    tot = sum(colw)
    for i, cw in enumerate(colw):
        tbl.columns[i].width = px(w * cw / tot)
    tbl.rows[0].height = px(head_h)
    for i in range(len(rows)):
        tbl.rows[i + 1].height = px(row_h)

    def fill_cell(cell, raw, bold, color, fill, sz, bottom, bottom_w):
        if fill is None:
            cell.fill.background()
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
        cell.margin_left = px(cell_pad)
        cell.margin_right = px(6)
        cell.margin_top = px(3)
        cell.margin_bottom = px(3)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = 1.22
        _fill_paragraph(p, str(raw) if str(raw).strip() else " ", sz, color, bold, FONT)
        p._p.get_or_add_endParaRPr().set("sz", str(int(sz * 100)))
        for side in "LRT":
            _cell_border(cell, side, None)
        _cell_border(cell, "B", bottom, bottom_w)

    for j, htxt in enumerate(head):
        hc = head_colors[j] if head_colors else head_color
        fill_cell(tbl.cell(0, j), htxt, True, hc, head_fill, size - 0.5, hc if head_fill is None else None, 1.25)
    for i, row in enumerate(rows):
        bg = PANEL if (zebra and i % 2 == 1) else None
        for j, cell in enumerate(row):
            col = (first_col_color or INK) if j == 0 else INK
            fill_cell(tbl.cell(i + 1, j), cell, first_col_bold and j == 0, col, bg, size, line, 0.75)
    return g


# ── 사진 ─────────────────────────────────────────────────────────────────────
def _load(p):
    if hasattr(p, "read"):
        p.seek(0)
    return ImageOps.exif_transpose(Image.open(p)).convert("RGB")


def _jpeg(im, q=90):
    b = io.BytesIO()
    im.save(b, "JPEG", quality=q, optimize=True)
    b.seek(0)
    return b


def crop_cover(p, w, h, cx=0.5, cy=0.5, max_w=1800):
    """사진을 w:h 비율로 꽉 채워 자른다(PIL). cx,cy 는 남길 중심(0~1). 파일 객체를 돌려준다."""
    im = _load(p)
    iw, ih = im.size
    s = max(w / iw, h / ih)
    cw, ch = w / s, h / s                      # 원본 좌표계에서의 잘라낼 크기
    x0 = (iw - cw) * cx
    y0 = (ih - ch) * cy
    im = im.crop((int(x0), int(y0), int(x0 + cw), int(y0 + ch)))
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    return _jpeg(im)


def crop_box(p, box, max_w=1800):
    """box = (l, t, r, b) 비율(0~1) 로 잘라낸 파일 객체."""
    im = _load(p)
    iw, ih = im.size
    l, t, r, b = box
    im = im.crop((int(l * iw), int(t * ih), int(r * iw), int(b * ih)))
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    return _jpeg(im)


def pic(sh_tree, root, rel, x, y, w, h, cap=None, cap_size=9, fit="contain", cap_h=None,
        border=False, cx=0.5, cy=0.5, cap_color=MUTED):
    """사진 한 장 + 설명. fit='contain' 은 상자 안에 다 들어가게(여백 생김), 'cover' 는 꽉 채우고 자른다."""
    p = pathlib.Path(root) / rel
    cap_h = cap_h if cap_h is not None else (24 if cap else 0)
    ah = h - cap_h - (5 if cap else 0)
    if fit == "contain":
        iw, ih = _load(p).size
        s = min(w / iw, ah / ih)
        dw, dh = iw * s, ih * s
        shp = sh_tree.add_picture(str(p), px(x + (w - dw) / 2), px(y + (ah - dh) / 2), px(dw), px(dh))
    else:
        shp = sh_tree.add_picture(crop_cover(p, w, ah, cx, cy), px(x), px(y), px(w), px(ah))
    if border:
        shp.line.color.rgb = LINE
        shp.line.width = Pt(0.75)
    if cap:
        text(sh_tree, x, y + ah + 5, w, cap_h, cap, size=cap_size, color=cap_color, spacing=1.3)
    return shp


# ── 영상 (PowerPoint 가 재생하는 임베드 + 슬라이드 열리면 자동 재생) ──────────
def video(sh_tree, mp4, poster, x, y, w, h):
    """H.264 MP4 를 슬라이드에 임베드한다. poster 는 재생 전(그리고 PDF)에 보이는 그림."""
    return sh_tree.add_movie(str(mp4), px(x), px(y), px(w), px(h), poster_frame_image=str(poster),
                             mime_type="video/mp4")


def autoplay(slide, movies, loop=False):
    """슬라이드가 열리면 movies 를 자동 재생한다 (PowerPoint 의 '시작: 자동으로' 와 같은 XML).

    PowerPoint 가 쓰는 구조를 그대로 만든다: 메인 시퀀스 안의 mediacall(playFrom) 효과 하나 +
    p:video/cMediaNode 하나가 영상마다 있다. 영상이 여럿이면 모두 동시에 시작한다.
    id 1 = 루트, 2 = 메인 시퀀스, 3 = 첫 단계(onBegin → 자동 시작), 4 = 묶음, 5~ = 영상별 노드.
    """
    from lxml import etree
    if not movies:
        return
    NS = ("xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\" "
          "xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\"")
    nxt = [5]

    def cid():
        v = nxt[0]
        nxt[0] += 1
        return v

    effects, media = [], []
    for mv in movies:
        spid = mv.shape_id
        effects.append(f"""
            <p:par><p:cTn id="{cid()}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect">
              <p:stCondLst><p:cond delay="0"/></p:stCondLst>
              <p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr>
                <p:cTn id="{cid()}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
              </p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>""")
    for mv in movies:
        spid = mv.shape_id
        rep = ' repeatCount="indefinite"' if loop else ""
        media.append(f"""
          <p:video><p:cMediaNode vol="80000"><p:cTn id="{cid()}" fill="hold" display="0"{rep}>
            <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
            <p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tn val="2"/></p:cond></p:endCondLst>
          </p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>""")
    xml = f"""<p:timing {NS}><p:tnLst><p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
        <p:seq concurrent="1" nextAc="seek">
          <p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
            <p:par><p:cTn id="3" fill="hold">
              <p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst>
              <p:childTnLst><p:par><p:cTn id="4" fill="hold">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                <p:childTnLst>{''.join(effects)}</p:childTnLst>
              </p:cTn></p:par></p:childTnLst>
            </p:cTn></p:par>
          </p:childTnLst></p:cTn>
          <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
          <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
        </p:seq>{''.join(media)}
      </p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""
    el = etree.fromstring(xml)
    sld = slide._element
    for old in sld.findall(qn("p:timing")):
        sld.remove(old)
    ext = sld.find(qn("p:extLst"))          # p:timing 은 p:clrMapOvr 뒤, p:extLst 앞
    if ext is not None:
        ext.addprevious(el)
    else:
        sld.append(el)
