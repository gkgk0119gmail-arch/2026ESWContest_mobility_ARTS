#!/usr/bin/env python3
"""결선 발표 대본(인쇄용) — present_script.py 의 쪽별 시간·시연 동작·말할 것을 A4 표로. 출력: …_결선발표_대본.docx (PDF 는 데스크톱 LibreOffice)."""
import pathlib
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from present_script import SCRIPT

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "2026ESWContest_모빌리티_ARTS_결선발표_대본.docx"
TITLES = {1: "표지", 2: "문제", 3: "해법", 4: "1차 방어", 5: "2차 방어", 6: "결과", 7: "안전 전용 RTOS", 8: "정리"}
doc = Document(); sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, m, Cm(1.5))
st = doc.styles["Normal"]; st.font.name = "NanumGothic"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "NanumGothic")

def para(t, size=10.5, bold=False, color=None, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); r = p.add_run(t); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = RGBColor(*color)
    return p

def shade(cell, hexcol):
    tc = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), hexcol); tc.append(s)

para("IcePredict 결선 발표 대본 — 10분 (발표 약 5분 + 시연 약 5분)", 16, True, (0, 44, 95), 6)
para("준비: 랜선(파이 ↔ 데스크톱 위쪽 포트) → 파이 터미널에서 sw/scripts/demo/venue_pi.sh (화면에 보이게) → 데스크톱 venue_desktop.sh (대기 화면, 발표 내내 켜 둔다: 1 노면 판정 · 2 1차 주행 · 3 2차 방어). "
     "아이패드 사진은 전체 화면, 밝기 최대. RealSense 창 '재촬영 보정'은 끔.", 10, False, (80, 80, 80), 8)
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
W = [Cm(2.6), Cm(2.0), Cm(7.2), Cm(14.9)]
for c, h in zip(t.rows[0].cells, ["쪽", "시간", "시연 동작", "말할 것"]):
    c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "002C5F")
for k in sorted(SCRIPT):
    tm, act, say = SCRIPT[k]; row = t.add_row().cells
    vals = [f"{k}. {TITLES[k]}", tm, act or "—", say]
    for i, (c, v) in enumerate(zip(row, vals)):
        c.text = ""; r = c.paragraphs[0].add_run(v); r.font.size = Pt(9.5 if i == 3 else 9)
        if i == 0: r.bold = True
        if i == 2 and act: r.font.color.rgb = RGBColor(190, 30, 30); r.bold = True
    if act: shade(row[0], "FDECEC")
for row in t.rows:
    for c, w in zip(row.cells, W): c.width = w
t.autofit = False
tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), W): gc.set(qn("w:w"), str(int(w.cm * 567)))   # LibreOffice 는 gridCol 너비를 본다
para("", after=2)
para("키: 1 노면 판정 · 2 1차 주행 · 3 2차 방어 · C 휠 영점 · D 2차 방어 켜기/끄기 · R 다시 · K 파이 메인 프로그램 죽이기 · I 빙판 표시 · T 3인칭 창 · (RealSense 창) +/− 판정 네모, B 재촬영 보정", 9.5, False, (80, 80, 80))
para("비상: 시연이 안 되면 제출 영상(IcePredict_시연영상_최종.mp4)을 튼다. 2차 방어가 '응답 없음'이면 랜선·venue_pi.sh 확인.", 9.5, False, (80, 80, 80))
doc.save(OUT); print("저장", OUT)
